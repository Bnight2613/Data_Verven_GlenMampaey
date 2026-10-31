"""
Data Verven - Glen Mampaey (Streamlit GUI)
==========================================
Run met:  streamlit run src/app.py

Dit is de GUI versie van de verf inventaris tool.
Hieronder staan de GUI componenten die je moet invullen.
"""

import streamlit as st
import pandas as pd
import openpyxl
import csv
import os
import io

# Pad naar de data bestanden
DATA_DIR = os.path.join(os.path.dirname(os.path.dirname(__file__)), "data")
PROJECTS_DIR = os.path.join(os.path.dirname(os.path.dirname(__file__)), "projects")
INVENTORY_CSV = os.path.join(DATA_DIR, "inventory.csv")


# ================================
# DEZE FUNCTIES KOMEN UIT main.py 
# ================================
def load_inventory(filepath=INVENTORY_CSV):
    paints = []
    with open(filepath, 'r', encoding='utf-8') as f:
        reader = csv.DictReader(f)
        for row in reader:
            row['aantal'] = int(row['aantal'])
            paints.append(row)
    return paints


def load_excel_inventory(uploaded_file):
    """Leest een Excel bestand en geeft de voorraad terug als lijst van dictionaries."""
    # Sla het geüploade bestand tijdelijk op zodat openpyxl het kan lezen
    content = uploaded_file.read()
    wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
    ws = wb['verfvoorraad']
    
    paints = []
    for row in ws.iter_rows(min_row=9, max_row=ws.max_row, values_only=True):
        # Kolommen: B(1)=Kleurnaam, C(2)=Artikelnummer, D(3)=Reeks, 
        # E(4)=Volledig nummer, F(5)=In voorraad, G(6)=Aantal
        if row[1] and row[3]:
            try:
                aantal = int(row[6]) if row[6] is not None else 0
            except (ValueError, TypeError):
                aantal = 0
            
            paints.append({
                'kleurnaam': str(row[1]).strip(),
                'artikelnummer': str(row[2]).strip() if row[2] else '',
                'reeks': str(row[3]).strip(),
                'volledig_nummer': str(row[4]).strip() if row[4] else '',
                'in_voorraad': str(row[5]).strip() if row[5] else 'Nee',
                'aantal': aantal,
                'bestellen': str(row[7]).strip() if row[7] else 'Nee',
                'aantal_bestellen': '0',
                'gewenst_aantal': '0'
            })
    
    return paints


def search_paint(paints, query):
    results = []
    query = query.lower()
    for p in paints:
        if (query in p['kleurnaam'].lower() or
            query in p['artikelnummer'].lower() or
            query in p['volledig_nummer'].lower()):
            results.append(p)
    return results


def shopping_list(paints, project_file):
    must_buy = []
    with open(project_file, 'r', encoding='utf-8') as f:
        reader = csv.DictReader(f)
        for item in reader:
            volledig = item['volledig_nummer']
            nodig = int(item['aantal_nodig'])
            gevonden = None
            for p in paints:
                if p['volledig_nummer'] == volledig:
                    gevonden = p
                    break
            if gevonden is None:
                must_buy.append({
                    'kleurnaam': item['kleurnaam'],
                    'volledig_nummer': volledig,
                    'status': 'NIET IN VOORRAAD',
                    'nodig': nodig,
                    'heeft': 0,
                    'tekort': nodig
                })
            elif gevonden['aantal'] < nodig:
                must_buy.append({
                    'kleurnaam': item['kleurnaam'],
                    'volledig_nummer': volledig,
                    'status': 'TE WEINIG',
                    'nodig': nodig,
                    'heeft': gevonden['aantal'],
                    'tekort': nodig - gevonden['aantal']
                })
    return must_buy


# ============================================================
# GUI CODE - Hier bouw jij de interface!
# ============================================================
# Handige Streamlit functies die je kan gebruiken:
#
#   st.title("Tekst")             → Grote titel
#   st.sidebar                    → Menu aan de linkerkant
#   st.text_input("Label")        → Tekst invoerveld
#   st.selectbox("Label", lijst)  → Dropdown menu
#   st.button("Tekst")            → Knop
#   st.dataframe(tabel)           → Mooie tabel tonen
#   st.metric("Label", waarde)    → Cijfer tonen
#   st.write("Tekst")             → Gewoon tekst tonen
#   st.info/warning()             → Gekleurde melding (blauw/oranje, kleurenblind-vriendelijk)
#   st.tabs(["Tab1", "Tab2"])     → Tabs maken


def main():
    # TODO 1: Zet een titel bovenaan
    st.title("Verfjes data")
    
    # Laad de voorraad (van CSV of van geüploade Excel)
    if 'uploaded_paints' in st.session_state:
        paints = st.session_state['uploaded_paints']
    else:
        paints = load_inventory()
    
    # TODO 3: Maak een zijbalk (sidebar) met een menu
    # Hint:
    st.sidebar.title("Menu")
    keuze = st.sidebar.radio("Kies:", ["Hoofd menu", "Zoeken", "Voorraad", "Bestellijst", "Excel uploaden"])
    # TODO 2: Toon stats bovenaan (hoeveel in voorraad, hoeveel niet)
    in_voorraad = len([p for p in paints if p['aantal'] > 0])
    niet_in_voorraad = len([p for p in paints if p['aantal'] == 0])
    if keuze == "Hoofd menu":
        col1, col2 = st.columns(2)
        col1.metric("In voorraad", in_voorraad)
        col2.metric("Niet in voorraad", niet_in_voorraad)
    
        st.write("### Per reeks:")
        reeksen = ["Model Air", "Model Color", "Metal Color", "Primer 17ml", "Primer 60ml", "Primer 200ml", "Enamel Revell"]
        for reeks in reeksen:
            aantal = len([p for p in paints if p['reeks'] == reeks and p['aantal'] > 0])
            st.write(f"- {reeks}: {aantal} in voorraad")


    # TODO 4: 
    if keuze == "Zoeken":
        query = st.text_input("Zoek verf: ")
        if query:
            results = search_paint(paints, query)
            st.write(f"gevonden: {len(results)} verven")
            st.dataframe(results)
    
    # TODO 5: 
    if keuze == "Voorraad":
        reeksen = ["Alle"] + list(set(p['reeks'] for p in paints))
        reeks_filter = st.selectbox("Filter op reeks:", reeksen)
    
        if reeks_filter == "Alle":
            in_voorraad = [p for p in paints if p['aantal'] > 0]
        else:
            in_voorraad = [p for p in paints if p['aantal'] > 0 and p['reeks'] == reeks_filter]
    
        st.dataframe(in_voorraad)

    
    # TODO 6:
    if keuze == "Bestellijst":
        bestand = st.file_uploader("Upload een project CSV:", type=["csv"])  # EERST dit!
    
        if bestand is not None:  # DAN pas checken
            content = bestand.read().decode('utf-8')
            reader = csv.DictReader(io.StringIO(content))
        
            must_buy = []
            for item in reader:
                volledig = item['volledig_nummer']
                nodig = int(item['aantal_nodig'])
            
                gevonden = None
                for p in paints:
                    if p['volledig_nummer'] == volledig:
                        gevonden = p
                        break
            
                if gevonden is None:
                    must_buy.append({
                        'kleurnaam': item['kleurnaam'],
                        'volledig_nummer': volledig,
                        'status': 'NIET IN VOORRAAD',
                        'nodig': nodig,
                        'heeft': 0,
                        'tekort': nodig
                    })
                elif gevonden['aantal'] < nodig:
                    must_buy.append({
                        'kleurnaam': item['kleurnaam'],
                        'volledig_nummer': volledig,
                        'status': 'TE WEINIG',
                        'nodig': nodig,
                        'heeft': gevonden['aantal'],
                        'tekort': nodig - gevonden['aantal']
                    })
        
        # dataframe PAS NA de loop!
            st.write(f"Te kopen: {len(must_buy)} verven")
            st.dataframe(must_buy)

            st.download_button("Download bestellijst", 
                pd.DataFrame(must_buy).to_csv(index=False),
                "bestellijst.csv")

                    

    # TODO 7: Excel uploaden om voorraad te updaten
    if keuze == "Excel uploaden":
        st.write("### Upload je Excel voorraad bestand")
        st.write("Upload het bestand `stock_verven_Datasheet.xlsx` om de voorraad in de app te updaten.")
        st.write("Dit werkt voor de huidige sessie. Sluit je de app, dan wordt de originele CSV weer gebruikt.")
        st.write("")
        
        excel_bestand = st.file_uploader("Kies je Excel bestand:", type=["xlsx"])
        
        if excel_bestand is not None:
            try:
                nieuwe_voorraad = load_excel_inventory(excel_bestand)
                st.session_state['uploaded_paints'] = nieuwe_voorraad
                st.info(f"✅ Voorraad geladen: {len(nieuwe_voorraad)} verven!")
                
                in_voorraad = len([p for p in nieuwe_voorraad if p['aantal'] > 0])
                niet_in_voorraad = len([p for p in nieuwe_voorraad if p['aantal'] == 0])
                
                col1, col2 = st.columns(2)
                col1.metric("In voorraad", in_voorraad)
                col2.metric("Niet in voorraad", niet_in_voorraad)
                
                st.write("Ga nu naar Zoeken, Voorraad of Bestellijst om de nieuwe data te gebruiken!")
            except Exception as e:
                st.warning(f"⚠️ Fout bij het lezen van het bestand: {e}")


if __name__ == "__main__":
    main()